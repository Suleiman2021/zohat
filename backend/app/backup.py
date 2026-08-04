"""نظام النسخ الاحتياطي — يرسل نسخة من قاعدة البيانات + ملف Excel لسجلات الشحنات
إلى تلغرام (أبسط طريقة مجانية موثوقة، بلا OAuth). يُشغَّل يدوياً أو بجدولة يومية."""
import io
import os
import sqlite3
import tempfile
import datetime as dt

import requests
from sqlmodel import Session, select
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .core import config
from .core.database import engine, db_file_path
from .models import Shipment

TG_API = "{base}/bot{token}/{method}"

# أعمدة تصدير سجل الشحنات (المفتاح، العنوان العربي)
_SHIP_COLS = [
    ("ref_no", "رقم القيد"), ("ship_date", "التاريخ"), ("created_by_name", "المسجِّل"),
    ("sender_name", "المرسِل"), ("receiver_name", "المستلِم"), ("driver_name", "السائق"),
    ("from_city", "جهة الإرسال"), ("to_city", "جهة الاستلام"),
    ("item_name", "الصنف"), ("item_code", "الكود"), ("brand", "الماركة"),
    ("origin_country", "المنشأ"), ("count", "العدد"), ("weight_kg", "الوزن (كغ)"),
    ("goods_value", "قيمة الفاتورة"), ("goods_price", "ثمن البضاعة"),
    ("financing", "التمويل"), ("bought_by", "من قام بالشراء"),
    ("fees_payment", "دفع الأجور"), ("export_status", "حالة التصدير"),
    ("export_date", "تاريخ التصدير"), ("delivery_status", "حالة التسليم"),
    ("collection_status", "حالة التحصيل"),
]


def shipments_xlsx() -> bytes:
    """يبني ملف Excel بكل سجلات الشحنات (بيانات خام من القاعدة)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "سجل الشحنات"
    ws.sheet_view.rightToLeft = True
    ws.append([label for _, label in _SHIP_COLS])
    fill = PatternFill("solid", fgColor="183C60")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    with Session(engine) as db:
        for s in db.exec(select(Shipment).order_by(Shipment.ref_no)).all():
            ws.append([_fmt(getattr(s, k, "")) for k, _ in _SHIP_COLS])
    for i, (_, label) in enumerate(_SHIP_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(label) + 3, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()[:10]
    return v


def db_snapshot() -> bytes | None:
    """لقطة متسقة من قاعدة SQLite عبر VACUUM INTO — آمنة حتى أثناء كتابة جارية،
    بخلاف قراءة الملف مباشرة التي قد تلتقط نسخة ناقصة. تعيد None إن لم تكن القاعدة SQLite."""
    path = db_file_path()
    if not path or not os.path.exists(path):
        return None
    fd, tmp = tempfile.mkstemp(suffix=".sqlite")
    os.close(fd)
    os.remove(tmp)                       # VACUUM INTO يرفض الكتابة فوق ملف موجود
    try:
        con = sqlite3.connect(path)
        try:
            con.execute("VACUUM INTO ?", (tmp,))
        finally:
            con.close()
        with open(tmp, "rb") as f:
            return f.read()
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _send_document(content: bytes, filename: str, caption: str = ""):
    url = TG_API.format(base=config.TELEGRAM_API_BASE.rstrip("/"),
                        token=config.TELEGRAM_BOT_TOKEN, method="sendDocument")
    r = requests.post(url, timeout=60,
                      data={"chat_id": config.TELEGRAM_CHAT_ID, "caption": caption},
                      files={"document": (filename, content)})
    r.raise_for_status()
    return r.json()


# ===================== جدولة تصمد أمام إعادة التشغيل =====================
# الخلل السابق: الخيط كان ينام 24 ساعة ثم يرسل، والعدّاد يبدأ من الصفر مع كل
# إعادة نشر — فلم يبلغ الموعد قط. الحل: مواعيد مثبَّتة بالساعة تُحفظ في القاعدة،
# فيُعرف ما أُرسل وما فات، ويُستدرك الفائت فور عودة الخدمة.
_K_SLOT, _K_OK, _K_ERR, _K_COUNT = ("backup_last_slot", "backup_last_ok",
                                    "backup_last_error", "backup_runs")


def local_now() -> dt.datetime:
    """الوقت الحالي بتوقيت دمشق (UTC+3 ثابت — لا توقيت صيفي في سوريا)."""
    return dt.datetime.utcnow() + dt.timedelta(hours=config.BACKUP_TZ_OFFSET)


def slot_key(now: dt.datetime, interval: int) -> str:
    """مفتاح آخر موعد مجدول انقضى، مثل «2026-08-05T03» لموعد الثالثة فجراً."""
    interval = max(1, int(interval))
    return f"{now:%Y-%m-%d}T{(now.hour // interval) * interval:02d}"


def next_slot_time(now: dt.datetime, interval: int) -> dt.datetime:
    """موعد الإرسال القادم بتوقيت دمشق."""
    interval = max(1, int(interval))
    base = now.replace(minute=0, second=0, microsecond=0)
    return base.replace(hour=(now.hour // interval) * interval) + dt.timedelta(hours=interval)


def _get(db, key: str, default: str = "") -> str:
    from .models import Setting
    row = db.get(Setting, key)
    return row.value if row else default


def _set(db, key: str, value: str):
    from .models import Setting
    row = db.get(Setting, key)
    if row:
        row.value = value
    else:
        row = Setting(key=key, value=value)
    db.add(row)


def backup_state() -> dict:
    """حالة الجدولة للعرض في الواجهة — تُثبت للمستخدم أن النظام يعمل فعلاً."""
    from sqlmodel import Session
    from .core.database import engine
    interval = config.BACKUP_INTERVAL_HOURS
    now = local_now()
    with Session(engine) as db:
        last_slot = _get(db, _K_SLOT)
        return {
            "enabled": interval > 0 and bool(config.TELEGRAM_BOT_TOKEN and config.TELEGRAM_CHAT_ID),
            "interval_hours": interval,
            "now_local": now.strftime("%Y-%m-%d %H:%M"),
            "last_slot": last_slot,
            "last_ok": _get(db, _K_OK),
            "last_error": _get(db, _K_ERR),
            "runs": int(_get(db, _K_COUNT, "0") or 0),
            "next_local": (next_slot_time(now, interval).strftime("%Y-%m-%d %H:%M")
                           if interval > 0 else ""),
        }


def _record(ok: bool, slot: str = "", error: str = ""):
    from sqlmodel import Session
    from .core.database import engine
    with Session(engine) as db:
        if ok:
            _set(db, _K_SLOT, slot)
            _set(db, _K_OK, local_now().strftime("%Y-%m-%d %H:%M"))
            _set(db, _K_ERR, "")
            _set(db, _K_COUNT, str(int(_get(db, _K_COUNT, "0") or 0) + 1))
        else:
            _set(db, _K_ERR, f"{local_now():%Y-%m-%d %H:%M} — {error}"[:400])
        db.commit()


def scheduler_loop(sleep=None):
    """يفحص كل دقيقة: هل انقضى موعد لم تُرسل نسخته بعد؟ فيرسلها.
    يستدرك المواعيد الفائتة (انقطاع/إعادة نشر) ولا يكرّر موعداً أُرسل."""
    import time
    sleep = sleep or time.sleep
    interval = config.BACKUP_INTERVAL_HOURS
    if interval <= 0:
        print("[backup] الجدولة معطَّلة (BACKUP_INTERVAL_HOURS = 0)")
        return
    if not (config.TELEGRAM_BOT_TOKEN and config.TELEGRAM_CHAT_ID):
        print("[!] [backup] تلغرام غير مضبوط — لن تُرسل نسخ تلقائية. "
              "اضبط TELEGRAM_BOT_TOKEN و TELEGRAM_CHAT_ID في متغيّرات البيئة.")
        return
    print(f"[ok] [backup] الجدولة تعمل: نسخة كل {interval} ساعة بتوقيت دمشق "
          f"(المواعيد: {', '.join(f'{h:02d}:00' for h in range(0, 24, max(1, interval)))})")
    while True:
        try:
            now = local_now()
            slot = slot_key(now, interval)
            from sqlmodel import Session
            from .core.database import engine
            with Session(engine) as db:
                done = _get(db, _K_SLOT)
            if done != slot:
                first = not done      # أول تشغيل: نرسل فوراً كإثبات أن الجدولة تعمل
                print(f"[backup] إرسال نسخة لموعد {slot}" + (" (أول تشغيل)" if first else ""))
                try:
                    run_backup()
                    _record(True, slot)
                    print(f"[ok] [backup] أُرسلت نسخة {slot} — الموعد القادم "
                          f"{next_slot_time(local_now(), interval):%Y-%m-%d %H:%M}")
                except Exception as e:
                    _record(False, error=str(e))
                    print(f"[!] [backup] فشل إرسال نسخة {slot}: {e}")
        except Exception as e:      # لا نسمح لأي خطأ بقتل الخيط
            print(f"[!] [backup] خطأ في حلقة الجدولة: {e}")
        sleep(60)


def run_backup() -> dict:
    """ينفّذ نسخة احتياطية كاملة → تلغرام. يرمي استثناءً عند الفشل."""
    if not (config.TELEGRAM_BOT_TOKEN and config.TELEGRAM_CHAT_ID):
        raise RuntimeError("لم تُضبط إعدادات تلغرام (TELEGRAM_BOT_TOKEN / TELEGRAM_CHAT_ID)")
    stamp = dt.datetime.now().strftime("%Y-%m-%d_%H%M")
    sent = []

    # 1) ملف Excel لسجلات الشحنات
    xlsx = shipments_xlsx()
    _send_document(xlsx, f"zohat_shipments_{stamp}.xlsx",
                   caption=f"📊 سجل الشحنات — {stamp}")
    sent.append("xlsx")

    # 2) نسخة من قاعدة البيانات — لقطة متسقة عبر VACUUM INTO
    #    (SQLite فقط؛ Postgres يُنسخ من Railway مباشرة)
    data = db_snapshot()
    if data:
        _send_document(data, f"zohat_db_{stamp}.sqlite",
                       caption=f"💾 نسخة قاعدة البيانات — {stamp}")
        sent.append("db")

    return {"ok": True, "sent": sent, "time": stamp}
