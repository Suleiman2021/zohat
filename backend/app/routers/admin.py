import io
import urllib.parse
from fastapi import APIRouter, Body, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..core.database import get_session
from ..core.security import admin_only, any_role, hash_pw
from ..models import User, Item

router = APIRouter(prefix="/api", tags=["admin"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.post("/export/xlsx")
def export_xlsx(payload: dict, user: User = Depends(any_role)):
    """تصدير أي جدول معروض إلى ملف Excel حقيقي (.xlsx) بترويسة منسّقة واتجاه من اليمين."""
    title = (payload.get("title") or "تقرير").strip()
    headers = payload.get("headers") or []
    rows = payload.get("rows") or []
    if not headers:
        raise HTTPException(400, "لا توجد أعمدة للتصدير")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title[:28] or "تقرير")
    ws.sheet_view.rightToLeft = True          # ورقة عربية من اليمين لليسار

    ws.append(list(headers))
    head_fill = PatternFill("solid", fgColor="183C60")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws.append(["" if c is None else c for c in row])

    # محاذاة وسط + عرض أعمدة تلقائي حسب أطول محتوى
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        longest = len(str(headers[col_idx - 1]))
        for row in rows:
            if col_idx <= len(row):
                longest = max(longest, len(str(row[col_idx - 1] if row[col_idx - 1] is not None else "")))
        ws.column_dimensions[letter].width = min(max(longest + 4, 10), 42)
        for cell in ws[letter][1:]:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"                    # تثبيت صف الرؤوس عند التمرير
    ws.auto_filter.ref = ws.dimensions        # فلاتر جاهزة على الأعمدة

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = urllib.parse.quote(f"{title}.xlsx")
    return StreamingResponse(buf, media_type=XLSX_MIME, headers={
        "Content-Disposition": f"attachment; filename=export.xlsx; filename*=UTF-8''{fname}"})


# ---- الأصناف (قاعدة البيانات) ----
@router.get("/items")
def items(db: Session = Depends(get_session), user: User = Depends(any_role),
          q: str = "", limit: int = 50, special: str = ""):
    """special: "1" = جدول 10% فقط، "0" = الأصناف العادية فقط، فارغ = الكل."""
    query = select(Item)
    if q:
        query = query.where(Item.name.contains(q) | Item.code.contains(q))
    if special == "1":
        query = query.where(Item.special_consumption == True)     # noqa: E712
    elif special == "0":
        query = query.where(Item.special_consumption == False)    # noqa: E712
    return db.exec(query.limit(min(limit, 5000))).all()


@router.post("/items/export")
def export_items(payload: dict, db: Session = Depends(get_session),
                 user: User = Depends(any_role)):
    """تصدير جدول الأصناف (العادي أو جدول 10%) إلى ملف Excel."""
    special = str(payload.get("special") or "")
    q = (payload.get("q") or "").strip()
    rows = items(db, user, q=q, limit=5000, special=special)
    title = "أصناف رسم الإنفاق 10%" if special == "1" else "الأصناف"
    return export_xlsx({"title": title,
                        "headers": ["الكود", "الصنف", "الرسم السوري للطن",
                                    "الرسم العراقي للطن", "جدول 10%"],
                        "rows": [[i.code, i.name, i.syrian_per_ton, i.iraqi_per_ton,
                                  "نعم" if i.special_consumption else "لا"] for i in rows]},
                       user)


@router.post("/items")
def add_item(it: Item, db: Session = Depends(get_session), user: User = Depends(admin_only)):
    db.add(it); db.commit(); db.refresh(it); return it


@router.put("/items/{iid}")
def edit_item(iid: int, patch: dict, db: Session = Depends(get_session),
              user: User = Depends(admin_only)):
    it = db.get(Item, iid)
    if not it: raise HTTPException(404)
    for k, v in patch.items():
        if hasattr(it, k): setattr(it, k, v)
    db.add(it); db.commit(); db.refresh(it); return it


@router.delete("/items/{iid}")
def delete_item(iid: int, db: Session = Depends(get_session), user: User = Depends(admin_only)):
    it = db.get(Item, iid)
    if it:
        db.delete(it); db.commit()
    return {"ok": True}


def _upsert_items(db: Session, rows: list[dict], special: bool = False) -> dict:
    """تحديث بالاسم إن وُجد، وإلا إضافة صنف جديد. مصدر مشترك لاستيراد CSV و xlsx.
    special: الاستيراد إلى «جدول 10%» — يُعلَّم به كل صنف مستورد."""
    existing = {i.name: i for i in db.exec(select(Item)).all()}
    added = updated = 0
    for row in rows:
        name = (row.get("name") or "").strip()
        if not name:
            continue
        rate = float(row.get("syrian_per_ton") or 0)
        iraqi = float(row.get("iraqi_per_ton") or 0)
        code = str(row.get("code") or "")
        if name in existing:
            existing[name].syrian_per_ton = rate
            existing[name].iraqi_per_ton = iraqi
            existing[name].code = code or existing[name].code
            if special:                     # الاستيراد لجدول 10% يعلّم الصنف
                existing[name].special_consumption = True
            db.add(existing[name]); updated += 1
        else:
            it = Item(code=code, name=name, syrian_per_ton=rate, iraqi_per_ton=iraqi,
                      special_consumption=special)
            db.add(it); existing[name] = it; added += 1
    db.commit()
    return {"added": added, "updated": updated}


@router.post("/items/bulk")
def bulk_items(payload: dict | list = Body(...), db: Session = Depends(get_session),
               user: User = Depends(admin_only)):
    """استيراد جماعي من CSV (مُحلَّل في الواجهة) — نفس منطق /items/import-xlsx.
    يقبل قائمة مباشرة (توافقاً) أو {rows, special}."""
    if isinstance(payload, list):
        return _upsert_items(db, payload)
    return _upsert_items(db, payload.get("rows") or [],
                         special=bool(payload.get("special")))


@router.post("/items/import-xlsx")
async def import_items_xlsx(file: UploadFile = File(...), special: str = "",
                            db: Session = Depends(get_session),
                            user: User = Depends(admin_only)):
    """استيراد مباشر من ملف إكسل «قاعدة البيانات» الأصلي (ورقة تحوي أعمدة الكود/الصنف/الرسم)."""
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb["قاعدة البيانات"] if "قاعدة البيانات" in wb.sheetnames else wb[wb.sheetnames[0]]

    # ابحث عن صف الرؤوس (يحوي عمود «الصنف») ضمن أول 10 صفوف، وحدّد مواقع الأعمدة بالاسم
    header_row = None
    col_map = {}
    for r in range(1, 11):
        row_vals = {c: ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)}
        for c, v in row_vals.items():
            if isinstance(v, str) and v.strip() == "الصنف":
                header_row = r
                for cc, vv in row_vals.items():
                    if not isinstance(vv, str):
                        continue
                    label = vv.strip()
                    if label == "الكود":
                        col_map["code"] = cc
                    elif label == "الصنف":
                        col_map["name"] = cc
                    elif "عراقي" in label:
                        col_map["iraqi_per_ton"] = cc
                    elif label.startswith("الرسم"):
                        col_map["syrian_per_ton"] = cc
                break
        if header_row:
            break
    if not header_row or "name" not in col_map:
        raise HTTPException(400, "لم يتم العثور على عمود «الصنف» في الملف — تأكد من ورقة قاعدة البيانات")

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=col_map["name"]).value
        if not name:
            continue
        code = ws.cell(row=r, column=col_map["code"]).value if "code" in col_map else ""
        rate = ws.cell(row=r, column=col_map["syrian_per_ton"]).value if "syrian_per_ton" in col_map else 0
        iraqi = ws.cell(row=r, column=col_map["iraqi_per_ton"]).value if "iraqi_per_ton" in col_map else 0
        rows.append({"code": str(code) if code is not None else "",
                     "name": str(name).strip(),
                     "syrian_per_ton": rate or 0,
                     "iraqi_per_ton": iraqi or 0})
    return _upsert_items(db, rows, special=(special == "1"))


# ---- المستخدمون (الإدارة فقط) ----
@router.get("/users")
def users(db: Session = Depends(get_session), user: User = Depends(admin_only)):
    return [{"id": u.id, "username": u.username, "full_name": u.full_name,
             "role": u.role, "branch": u.branch, "is_active": u.is_active}
            for u in db.exec(select(User)).all()]


@router.post("/users")
def add_user(payload: dict, db: Session = Depends(get_session), user: User = Depends(admin_only)):
    if db.exec(select(User).where(User.username == payload["username"])).first():
        raise HTTPException(400, "اسم المستخدم موجود")
    u = User(username=payload["username"], full_name=payload.get("full_name", ""),
             role=payload.get("role", "branch"), branch=payload.get("branch", ""),
             hashed_password=hash_pw(payload["password"]))
    db.add(u); db.commit(); db.refresh(u)
    return {"id": u.id, "username": u.username}


@router.put("/users/{uid}")
def edit_user(uid: int, payload: dict, db: Session = Depends(get_session),
              user: User = Depends(admin_only)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "المستخدم غير موجود")
    # تغيير اسم المستخدم مع فحص التكرار
    new_username = (payload.get("username") or "").strip()
    if new_username and new_username != u.username:
        if db.exec(select(User).where(User.username == new_username)).first():
            raise HTTPException(400, "اسم المستخدم موجود")
        u.username = new_username
    if "full_name" in payload: u.full_name = payload["full_name"]
    if payload.get("role"): u.role = payload["role"]
    if "branch" in payload: u.branch = payload["branch"]
    if "is_active" in payload: u.is_active = bool(payload["is_active"])
    if payload.get("password"):
        u.hashed_password = hash_pw(payload["password"])
    # لا نترك النظام دون أي مدير فعّال
    if u.role != "admin" or not u.is_active:
        others = db.exec(select(User).where(User.role == "admin", User.is_active == True,
                                            User.id != uid)).all()
        if not others:
            raise HTTPException(400, "لا يمكن ترك النظام دون مدير فعّال")
    db.add(u); db.commit(); db.refresh(u)
    return {"id": u.id, "username": u.username}


@router.delete("/users/{uid}")
def delete_user(uid: int, db: Session = Depends(get_session), user: User = Depends(admin_only)):
    u = db.get(User, uid)
    if not u:
        return {"ok": True}
    if u.id == user.id:
        raise HTTPException(400, "لا يمكنك حذف حسابك الحالي")
    if u.role == "admin":
        others = db.exec(select(User).where(User.role == "admin", User.id != uid)).all()
        if not others:
            raise HTTPException(400, "لا يمكن حذف آخر مدير في النظام")
    db.delete(u); db.commit()
    return {"ok": True}
